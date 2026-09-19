"""Page Markdown 원고를 한글 폰트가 포함된 Word 문서로 변환한다."""
from __future__ import annotations

import json
import re
import struct
import uuid
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

BASE = Path(__file__).resolve().parents[1]
OUT = BASE / 'output' / 'PlaceOS_Page_논문_20260919.docx'
FONT = 'Malgun Gothic'
NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def font_run(run, size: float | None = None) -> None:
    run.font.name = FONT
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size:
        run.font.size = Pt(size)
    rf = run._element.get_or_add_rPr().rFonts
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + key), FONT)


def hyperlink(paragraph, text: str, url: str | None = None, anchor: str | None = None) -> None:
    link = OxmlElement('w:hyperlink')
    if url:
        rid = paragraph.part.relate_to(url, RNS + '/hyperlink', is_external=True)
        link.set(qn('r:id'), rid)
    if anchor:
        link.set(qn('w:anchor'), anchor)
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), FONT)
    pr.append(fonts)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '333333'); pr.append(color)
    if anchor:
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '17'); pr.append(sz)
    r.append(pr)
    t = OxmlElement('w:t'); t.text = text; r.append(t)
    link.append(r); paragraph._p.append(link)


def embed_fonts(path: Path) -> None:
    with ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    table = etree.fromstring(parts['word/fontTable.xml'])
    relpath = 'word/_rels/fontTable.xml.rels'
    rels = etree.fromstring(parts[relpath]) if relpath in parts else etree.Element('Relationships', nsmap={None: REL})
    f = next((x for x in table if x.get(qn('w:name')) == FONT), None)
    if f is None:
        f = etree.SubElement(table, qn('w:font')); f.set(qn('w:name'), FONT)
    for label, filename in [('Regular', 'malgun.ttf'), ('Bold', 'malgunbd.ttf')]:
        fp = Path('C:/Windows/Fonts') / filename
        data = bytearray(fp.read_bytes())
        tables = struct.unpack_from('>H', data, 4)[0]
        os2 = next(struct.unpack_from('>I', data, 12+i*16+8)[0] for i in range(tables) if data[12+i*16:16+i*16] == b'OS/2')
        assert not struct.unpack_from('>H', data, os2+8)[0] & 2, '폰트 임베딩 제한'
        key = uuid.uuid4(); mask = bytes.fromhex(key.hex)[::-1]
        for i in range(32):
            data[i] ^= mask[i % 16]
        name = f'page-{label.lower()}.odttf'
        parts['word/fonts/' + name] = bytes(data)
        rid = 'rIdPage' + label
        etree.SubElement(rels, '{' + REL + '}Relationship', Id=rid, Type=RNS+'/font', Target='fonts/'+name)
        node = etree.SubElement(f, qn('w:embed' + label))
        node.set(qn('r:id'), rid); node.set(qn('w:fontKey'), '{'+str(key).upper()+'}')
    types = etree.fromstring(parts['[Content_Types].xml'])
    if not any(x.get('Extension') == 'odttf' for x in types):
        etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Default', Extension='odttf', ContentType='application/vnd.openxmlformats-officedocument.obfuscatedFont')
    for name, element in [('word/fontTable.xml', table), (relpath, rels), ('[Content_Types].xml', types)]:
        parts[name] = etree.tostring(element, xml_declaration=True, encoding='UTF-8', standalone=True)
    with ZipFile(path, 'w', ZIP_DEFLATED) as z:
        for name, content in parts.items():
            z.writestr(name, content)


def main() -> None:
    source = (BASE/'paper-page.md').read_text(encoding='utf-8')
    source = re.sub(r'<!--.*?-->', '', source, flags=re.S)
    source = '## 초록' + source.split('## 초록', 1)[1]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.1); sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    sec.footer_distance = Cm(1.0)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
        style = doc.styles[name]
        style.font.name = FONT; style.font.color.rgb = RGBColor(0,0,0)
        rf = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ('ascii','hAnsi','eastAsia','cs'):
            rf.set(qn('w:'+key),FONT)
        for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            rf.attrib.pop(qn('w:'+key),None)
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for name,size in [('Title',19),('Heading 1',14),('Heading 2',11.5)]:
        st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True
        st.paragraph_format.space_before=Pt(14 if name!='Title' else 0)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.keep_with_next=True
    title=doc.add_paragraph('공공행정자료 결합 기반 상업용 공실 정보의\n데이터 품질과 계산 재현성', 'Title')
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    sub=doc.add_paragraph('PlaceOS 사례', 'Subtitle'); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.styles['Subtitle'].font.size=Pt(12)
    doc.styles['Subtitle'].paragraph_format.space_after=Pt(14)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
    for name in ('embedTrueTypeFonts','saveSubsetFonts'):
        node=OxmlElement('w:'+name); node.set(qn('w:val'),'true' if name=='embedTrueTypeFonts' else 'false'); doc.settings.element.append(node)
    doc.core_properties.title='공공행정자료 결합 기반 상업용 공실 정보의 데이터 품질과 계산 재현성'
    doc.core_properties.subject='PlaceOS Page 연구'
    doc.core_properties.author=''; doc.core_properties.last_modified_by=''
    notes: dict[str, tuple[int,str]]={}

    def inline(p, text):
        pattern=r'(\[[^\]]+\]\([^)]+\)|\*\*.*?\*\*|`[^`]+`|\*[^*]+\*|https?://[^\s]+)'
        for token in re.split(pattern,text):
            if not token: continue
            m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
            if m:
                label,target=m.groups()
                if target.startswith('http'):
                    hyperlink(p,label,target)
                else:
                    if target not in notes: notes[target]=(len(notes)+1,label)
                    n=notes[target][0]; hyperlink(p,f' [자료 {n}]',anchor=f'source{n}')
            elif token.startswith('http'):
                hyperlink(p,token.rstrip('.'),token.rstrip('.'))
            else:
                bold=token.startswith('**'); italic=token.startswith('*') and not bold
                value=token.strip('*') if bold or italic else token.strip('`')
                r=p.add_run(value); font_run(r); r.bold=bold; r.italic=italic

    lines=source.splitlines(); i=0
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line or line=='---': continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip()); i+=1
            rows=[r for r in rows if not re.match(r'^\|[\s:|\-]+\|$',r)]
            values=[[x.strip() for x in r.strip('|').split('|')] for r in rows]
            doc.add_paragraph('표 1 분석 단위와 주요 변수',style='Caption')
            table=doc.add_table(rows=0,cols=len(values[0])); table.autofit=False
            widths=[2.5,5.1,9.0]
            for col,w in zip(table.columns,widths): col.width=Cm(w)
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            table._tbl.tblPr.append(borders)
            for j,row in enumerate(values):
                cells=table.add_row().cells
                trpr=table.rows[-1]._tr.get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
                if j==0: trpr.append(OxmlElement('w:tblHeader'))
                for k,(cell,value) in enumerate(zip(cells,row)):
                    cell.width=Cm(widths[k]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pr=cell._tc.get_or_add_tcPr()
                    margins=OxmlElement('w:tcMar')
                    for side in ('top','bottom','left','right'):
                        el=OxmlElement('w:'+side); el.set(qn('w:w'),'100'); el.set(qn('w:type'),'dxa'); margins.append(el)
                    pr.append(margins)
                    if j==0:
                        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EDF2'); pr.append(shade)
                    p=cell.paragraphs[0]; p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
                    if k==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    inline(p,value)
                    for r in p.runs: font_run(r,9); r.bold=j==0
            continue
        if line.startswith('#'):
            level=1 if line.startswith('## ') else 2
            heading=re.sub(r'^#+\s*','',line)
            heading=re.sub(r'^(\d+)\. ',r'\1 ',heading)
            heading=heading.replace('·',' ').replace('—',' ').replace('의 평가 절차','의 평가 절차')
            p=doc.add_paragraph(heading,style=f'Heading {level}')
            continue
        p=doc.add_paragraph(); inline(p,line)
    doc.add_heading('자료 주석',level=1)
    doc.add_paragraph('다음 자료는 원고의 검사 입력과 기존 실행 결과를 연결한다. 경로는 저장소의 docs/papers를 기준으로 하며, 파일 목록이 원천 전체의 공개 접근을 보장하지는 않는다.')
    for target,(n,label) in notes.items():
        p=doc.add_paragraph()
        start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(n)); start.set(qn('w:name'),f'source{n}'); p._p.append(start)
        r=p.add_run(f'{n}. {label}'); font_run(r,9); r.bold=True
        end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(n)); p._p.append(end)
        p.add_run('\n'); hyperlink(p,target,'../'+target)
        p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(6)
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT); embed_fonts(OUT)
    print(json.dumps({'docx':str(OUT),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'source_notes':len(notes),'embedded_fonts':2},ensure_ascii=False))


if __name__=='__main__':
    main()
