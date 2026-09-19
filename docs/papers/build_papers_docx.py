"""네 논문의 Markdown을 같은 판형의 Word 원고로 만든다."""
from pathlib import Path
import importlib.util
import json
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location('page_builder', BASE/'page-study/build_page_docx.py')
helper = importlib.util.module_from_spec(spec)
spec.loader.exec_module(helper)
TITLES = {
    'Platform': '거점 사전분포를 고려한 상권 그래프 신경망 업종 분류의 성능 한계',
    'Page': '공공행정자료 결합 기반 상업용 공실 정보의 데이터 품질과 계산 재현성',
    'Posting': '공개 데이터 기반 상가 입점 회수기간 시나리오의 해상도 한계',
    'Program': '창업 검증 프로그램 생성의 근거 계약과 사실성 가드레일',
}

def build(track, title):
    src = (BASE/f'paper-{track.lower()}.md').read_text(encoding='utf-8-sig')
    src = re.sub(r'<!--.*?-->', '', src, flags=re.S)
    src = '## 초록' + src.split('## 초록', 1)[1]
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(2.1), Cm(2)
    s.left_margin = s.right_margin = Cm(2.2)
    s.footer_distance = Cm(1)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','Footer'):
        st = doc.styles[name]
        st.font.name = helper.FONT
        st.font.color.rgb = RGBColor(0,0,0)
        rf = st.element.get_or_add_rPr().get_or_add_rFonts()
        for k in ('ascii','hAnsi','eastAsia','cs'): rf.set(qn('w:'+k), helper.FONT)
        for k in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'): rf.attrib.pop(qn('w:'+k),None)
    n = doc.styles['Normal']
    n.font.size = Pt(10.5)
    n.paragraph_format.line_spacing = 1.4
    n.paragraph_format.space_after = Pt(7)
    n.paragraph_format.widow_control = True
    for name,size in [('Title',19),('Heading 1',14),('Heading 2',11.5)]:
        st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True
        st.paragraph_format.space_before=Pt(13 if name!='Title' else 0)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.keep_with_next=True
    p=doc.add_paragraph(title,'Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('PlaceOS '+track,'Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.styles['Subtitle'].font.size=Pt(11)
    doc.styles['Subtitle'].paragraph_format.space_after=Pt(14)
    p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
    doc.core_properties.title=title
    doc.core_properties.subject='PlaceOS '+track+' 연구'
    doc.core_properties.author=''; doc.core_properties.last_modified_by=''
    for name in ('embedTrueTypeFonts','saveSubsetFonts'):
        el=OxmlElement('w:'+name); el.set(qn('w:val'),'true' if name=='embedTrueTypeFonts' else 'false'); doc.settings.element.append(el)
    notes={}

    def inline(p, text):
        pattern=r'(\[[^\]]+\]\([^)]+\)|\*\*.*?\*\*|`[^`]+`|\*[^*]+\*|https?://[^\s]+)'
        for tok in re.split(pattern,text):
            if not tok: continue
            m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',tok)
            if m:
                label,target=m.groups()
                if target.startswith('http'): helper.hyperlink(p,label,target)
                else:
                    if target not in notes: notes[target]=(len(notes)+1,label)
                    helper.hyperlink(p,label+f' [자료 {notes[target][0]}]',anchor=f'source{notes[target][0]}')
            elif tok.startswith('http'): helper.hyperlink(p,tok.rstrip('.'),tok.rstrip('.'))
            else:
                bold=tok.startswith('**'); italic=tok.startswith('*') and not bold
                r=p.add_run(tok.strip('*') if bold or italic else tok.strip('`'))
                helper.font_run(r); r.bold=bold; r.italic=italic

    lines=src.splitlines(); i=0; table_count=0
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line or line=='---' or line.startswith('<a '): continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip()); i+=1
            rows=[x for x in rows if not re.match(r'^\|[\s:|\-]+\|$',x)]
            values=[[v.strip() for v in x.strip('|').split('|')] for x in rows]
            count=len(values[0]); table_count+=1
            assert all(len(r)==count for r in values)
            widths = ([2.5,5.1,9.0] if track=='Page' and count==3 else
                      [4.6,4,8] if count==3 else
                      [2,2,4.6,8] if count==4 else
                      [4.2]+[12.4/(count-1)]*(count-1))
            t=doc.add_table(rows=0,cols=count); t.autofit=False
            for col,w in zip(t.columns,widths): col.width=Cm(w)
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            t._tbl.tblPr.append(borders)
            for j,row in enumerate(values):
                cells=t.add_row().cells
                trpr=t.rows[-1]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
                if j==0: trpr.append(OxmlElement('w:tblHeader'))
                for k,(cell,value) in enumerate(zip(cells,row)):
                    cell.width=Cm(widths[k]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pr=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
                    for side in ('top','bottom','left','right'):
                        el=OxmlElement('w:'+side); el.set(qn('w:w'),'100'); el.set(qn('w:type'),'dxa'); margins.append(el)
                    pr.append(margins)
                    if j==0:
                        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EDF2'); pr.append(shade)
                    p=cell.paragraphs[0]; p.paragraph_format.line_spacing=1.15
                    p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
                    if len(value)<20: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    inline(p,value)
                    for r in p.runs: helper.font_run(r,9); r.bold=j==0
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
        elif line.startswith('#'):
            heading=re.sub(r'^#+\s*','',line)
            heading=re.sub(r'[^\w\s가-힣]',' ',heading)
            heading=re.sub(r'\s+',' ',heading).strip()
            doc.add_paragraph(heading,style='Heading 1' if line.startswith('## ') else 'Heading 2')
        else:
            while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||---|[-*] |\d+\. )',lines[i].strip()):
                line+=' '+lines[i].strip(); i+=1
            p=doc.add_paragraph()
            if line.startswith('- '): line='• '+line[2:]
            inline(p,line)
    if notes:
        doc.add_heading('자료 주석',level=1)
        doc.add_paragraph('자료 경로는 저장소의 docs/papers를 기준으로 한다. 원천 전체의 공개 접근이나 독립 재현을 보장하지 않는다.')
        for target,(num,label) in notes.items():
            p=doc.add_paragraph()
            el=OxmlElement('w:bookmarkStart'); el.set(qn('w:id'),str(num)); el.set(qn('w:name'),f'source{num}'); p._p.append(el)
            r=p.add_run(f'{num}. {label}'); helper.font_run(r,9); r.bold=True
            el=OxmlElement('w:bookmarkEnd'); el.set(qn('w:id'),str(num)); p._p.append(el)
            r=p.add_run('\n'+target); helper.font_run(r,9)
            p.paragraph_format.line_spacing=1.15
    out=BASE/'output'/f'PlaceOS_{track}_논문_20260919.docx'
    out.parent.mkdir(exist_ok=True)
    doc.save(out); helper.embed_fonts(out)
    return {'track':track,'path':str(out),'paragraphs':len(doc.paragraphs),'tables':table_count}

if __name__=='__main__':
    print(json.dumps([build(k,v) for k,v in TITLES.items()],ensure_ascii=False,indent=2))
